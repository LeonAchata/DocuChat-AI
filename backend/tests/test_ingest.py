from pathlib import Path

import pytest

from docuchat.ingest.chunker import chunk_blocks, count_tokens
from docuchat.ingest.parsers import Block, UnsupportedFormatError, parse_file
from docuchat.service import Components
from tests.conftest import HANDBOOK, make_components


def test_markdown_keeps_heading_structure(tmp_path: Path) -> None:
    path = tmp_path / "handbook.md"
    path.write_text(HANDBOOK, encoding="utf-8")
    doc = parse_file(path)
    assert doc.title == "Employee Handbook"
    headings = [(b.heading_level, b.text) for b in doc.blocks if b.heading_level]
    assert headings[:2] == [(1, "Employee Handbook"), (2, "Vacation policy")]


def test_html_strips_boilerplate(tmp_path: Path) -> None:
    path = tmp_path / "page.html"
    path.write_text(
        "<html><head><title>Guide</title><script>x=1</script></head><body><nav>menu</nav>"
        "<h1>Intro</h1><p>Hello   world</p></body></html>",
        encoding="utf-8",
    )
    doc = parse_file(path)
    assert doc.title == "Guide"
    assert [b.text for b in doc.blocks] == ["Intro", "Hello world"]


def test_docx_headings_and_tables(tmp_path: Path) -> None:
    import docx

    document = docx.Document()
    document.add_heading("Quarterly report", level=1)
    document.add_paragraph("Revenue grew 12 percent.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Q1", "10M"
    path = tmp_path / "report.docx"
    document.save(str(path))

    doc = parse_file(path)
    assert doc.blocks[0].heading_level == 1
    assert any("Q1 | 10M" in b.text for b in doc.blocks)


def test_pdf_pages_and_font_size_headings(tmp_path: Path) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    path = tmp_path / "paper.pdf"
    pdf = canvas.Canvas(str(path), pagesize=A4)
    for page, (heading, body) in enumerate(
        [
            ("Introduction", "Transformers changed NLP."),
            ("Results", "Accuracy reached 91 percent."),
        ],
        start=1,
    ):
        pdf.setFont("Helvetica-Bold", 20)
        pdf.drawString(72, 760, heading)
        pdf.setFont("Helvetica", 11)
        for i in range(4):
            pdf.drawString(72, 720 - i * 14, f"{body} Line {i} of page {page}")
        pdf.showPage()
    pdf.save()

    doc = parse_file(path)
    assert doc.pages == 2
    headings = [(b.text, b.page) for b in doc.blocks if b.heading_level]
    assert headings == [("Introduction", 1), ("Results", 2)]
    assert any(b.page == 2 and "91 percent" in b.text for b in doc.blocks)


def test_unsupported_format(tmp_path: Path) -> None:
    path = tmp_path / "image.png"
    path.write_bytes(b"\x89PNG")
    with pytest.raises(UnsupportedFormatError):
        parse_file(path)


def test_chunks_respect_sections_budget_and_overlap() -> None:
    sentence = "The quick brown fox jumps over the lazy dog near the river bank. "
    blocks = [
        Block("Guide", heading_level=1),
        Block("Setup", heading_level=2),
        Block(sentence * 12, page=1),
        Block("Usage", heading_level=2),
        Block("Run the tool with the default options.", page=2),
    ]
    chunks = chunk_blocks(blocks, max_tokens=40, overlap_tokens=15)
    setup = [c for c in chunks if c.section == "Guide > Setup"]
    usage = [c for c in chunks if c.section == "Guide > Usage"]
    assert len(setup) > 2 and len(usage) == 1
    assert all(c.tokens <= 40 for c in chunks)
    assert usage[0].text == "Run the tool with the default options."  # no overlap across sections
    first_tail = setup[0].text.split(". ")[-1]
    assert first_tail.rstrip(".") in setup[1].text  # sentence overlap between neighbours
    assert setup[0].page_start == 1 and usage[0].page_start == 2


def test_token_estimate() -> None:
    assert count_tokens("Hello, world!") == 4


def test_ingest_is_idempotent_and_indexes_breadcrumbs(
    components: Components, docs_dir: Path
) -> None:
    docs = components.store.list_documents()
    assert {d.status for d in docs} == {"ready"}
    again = components.ingestor.ingest(docs_dir / "handbook.md")
    assert len(components.store.list_documents()) == 2
    assert again.id in {d.id for d in docs}

    hits = components.store.lexical_search("remote stipend", 1)
    chunk = hits[0][0]
    assert chunk.section == "Employee Handbook > Remote work"
    assert chunk.search_text.startswith("Employee Handbook > Employee Handbook > Remote work")


def test_failed_ingestion_is_recorded(components: Components, tmp_path: Path) -> None:
    empty = tmp_path / "empty.txt"
    empty.write_text("   \n", encoding="utf-8")
    doc = components.ingestor.ingest(empty)
    assert doc.status == "failed"
    assert "No extractable text" in (doc.error or "")


def test_contextualizer_feeds_search_text(settings, docs_dir: Path) -> None:  # type: ignore[no-untyped-def]
    comps = make_components(settings)
    calls: list[int] = []

    def contextualize(document: str, chunks: list[str]) -> list[str]:
        calls.append(len(chunks))
        return [f"CONTEXT-{i}" for i in range(len(chunks))]

    comps.ingestor.contextualizer = contextualize
    comps.ingestor.ingest(docs_dir / "security.md")
    chunk = comps.store.lexical_search("CONTEXT", 1)[0][0]
    assert "CONTEXT-" in chunk.search_text and "CONTEXT" not in chunk.text
    assert calls


def test_markdown_reflows_prose_but_keeps_lists(tmp_path: Path) -> None:
    path = tmp_path / "notes.md"
    path.write_text("A wrapped\nparagraph here.\n\n- item one\n- item two\n", encoding="utf-8")
    blocks = parse_file(path).blocks
    assert blocks[0].text == "A wrapped paragraph here."
    assert blocks[1].text == "- item one\n- item two"
