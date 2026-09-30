"""Turn files into a flat list of blocks that keep the structure chunking needs:
the page each block came from and whether it is a heading (and at which level)."""

import re
import statistics
from dataclasses import dataclass
from pathlib import Path
from typing import Any

_SPACES_RE = re.compile("[ \\t\\u00a0]+")  # spaces, tabs, non-breaking spaces

_STRUCTURED_LINE = re.compile(r"^\s*([-*+]\s|\d+[.)]\s|\||```|    )")  # lists, tables, code

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".md", ".markdown", ".txt", ".html", ".htm"}


@dataclass
class Block:
    text: str
    page: int | None = None
    heading_level: int | None = None  # 1 = top level; None = body text


@dataclass
class ParsedDocument:
    title: str
    blocks: list[Block]
    pages: int | None = None


class UnsupportedFormatError(ValueError):
    pass


def parse_file(path: Path, *, title: str | None = None) -> ParsedDocument:
    ext = path.suffix.lower()
    if ext == ".pdf":
        doc = _parse_pdf(path)
    elif ext == ".docx":
        doc = _parse_docx(path)
    elif ext in (".md", ".markdown"):
        doc = _parse_markdown(path.read_text(encoding="utf-8", errors="replace"))
    elif ext in (".html", ".htm"):
        doc = _parse_html(path.read_text(encoding="utf-8", errors="replace"))
    elif ext == ".txt":
        doc = _parse_text(path.read_text(encoding="utf-8", errors="replace"))
    else:
        raise UnsupportedFormatError(f"Unsupported file type: {ext or 'no extension'}")
    doc.title = title or doc.title or path.stem.replace("_", " ").replace("-", " ")
    doc.blocks = [b for b in doc.blocks if b.text.strip()]
    return doc


# ---- PDF -------------------------------------------------------------------------------------


def _parse_pdf(path: Path) -> ParsedDocument:
    import pdfplumber

    blocks: list[Block] = []
    title = ""
    with pdfplumber.open(path) as pdf:
        meta_title = (pdf.metadata or {}).get("Title")
        lines_by_page = [
            (number, page.extract_text_lines(keep_blank_chars=False))
            for number, page in enumerate(pdf.pages, start=1)
        ]
        page_count = len(pdf.pages)

    sizes = [_line_size(line) for _, lines in lines_by_page for line in lines]
    body_size = statistics.median(sizes) if sizes else 0.0
    heading_sizes = sorted({round(s, 1) for s in sizes if s > body_size * 1.15}, reverse=True)

    for number, lines in lines_by_page:
        paragraph: list[str] = []
        prev_bottom: float | None = None
        for line in lines:
            text = _clean(line["text"])
            if not text:
                continue
            size = _line_size(line)
            level = _heading_level(text, size, heading_sizes)
            gap = (line["top"] - prev_bottom) if prev_bottom is not None else 0
            prev_bottom = line["bottom"]
            if level is not None:
                _flush(paragraph, blocks, number)
                blocks.append(Block(text, number, level))
                if not title and level == 1:
                    title = text
                continue
            if paragraph and gap > size * 0.9:  # vertical whitespace => new paragraph
                _flush(paragraph, blocks, number)
            paragraph.append(text)
        _flush(paragraph, blocks, number)

    return ParsedDocument(
        title=meta_title or title, blocks=_merge_hyphenation(blocks), pages=page_count
    )


def _line_size(line: dict[str, Any]) -> float:
    chars = line.get("chars") or []
    return statistics.median(c["size"] for c in chars) if chars else 0.0


def _heading_level(text: str, size: float, heading_sizes: list[float]) -> int | None:
    if not heading_sizes or len(text) > 120 or text.endswith((".", ",", ";")):
        return None
    rounded = round(size, 1)
    if rounded not in heading_sizes:
        return None
    return min(heading_sizes.index(rounded) + 1, 6)


def _flush(paragraph: list[str], blocks: list[Block], page: int) -> None:
    if paragraph:
        blocks.append(Block(" ".join(paragraph), page))
        paragraph.clear()


def _merge_hyphenation(blocks: list[Block]) -> list[Block]:
    for block in blocks:
        block.text = re.sub(r"(\w)- (\w)", r"\1\2", block.text)
    return blocks


# ---- DOCX ------------------------------------------------------------------------------------


def _parse_docx(path: Path) -> ParsedDocument:
    import docx

    document = docx.Document(str(path))
    blocks: list[Block] = []
    for para in document.paragraphs:
        text = _clean(para.text)
        if not text:
            continue
        style = (para.style.name if para.style is not None else "") or ""
        level = None
        if style.lower() == "title":
            level = 1
        elif match := re.match(r"heading (\d)", style.lower()):
            level = int(match.group(1))
        blocks.append(Block(text, None, level))
    for table in document.tables:
        rows = [" | ".join(_clean(c.text) for c in row.cells) for row in table.rows]
        blocks.append(Block("\n".join(rows)))
    title = document.core_properties.title or next(
        (b.text for b in blocks if b.heading_level == 1), ""
    )
    return ParsedDocument(title=title, blocks=blocks)


# ---- Markdown / HTML / text ------------------------------------------------------------------


def _parse_markdown(source: str) -> ParsedDocument:
    blocks: list[Block] = []
    paragraph: list[str] = []
    in_code = False

    def flush() -> None:
        if paragraph:
            structured = any(_STRUCTURED_LINE.match(line) for line in paragraph)
            # Re-flow hard-wrapped prose; keep line breaks for lists, tables and code.
            text = "\n".join(paragraph) if structured else " ".join(p.strip() for p in paragraph)
            blocks.append(Block(text))
            paragraph.clear()

    for raw in source.splitlines():
        line = raw.rstrip()
        if line.lstrip().startswith("```"):
            in_code = not in_code
            paragraph.append(line)
            continue
        heading = None if in_code else re.match(r"^(#{1,6})\s+(.*)", line)
        if heading:
            flush()
            blocks.append(Block(heading.group(2).strip(" #"), None, len(heading.group(1))))
        elif not line.strip() and not in_code:
            flush()
        else:
            paragraph.append(line)
    flush()
    title = next((b.text for b in blocks if b.heading_level == 1), "")
    return ParsedDocument(title=title, blocks=blocks)


def _parse_html(source: str) -> ParsedDocument:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(source, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header", "noscript"]):
        tag.decompose()
    blocks: list[Block] = []
    for el in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "pre", "td"]):
        text = _clean(el.get_text(" "))
        if text:
            level = int(el.name[1]) if el.name.startswith("h") and len(el.name) == 2 else None
            blocks.append(Block(text, None, level))
    title = _clean(soup.title.get_text()) if soup.title else ""
    return ParsedDocument(title=title, blocks=blocks)


def _parse_text(source: str) -> ParsedDocument:
    paragraphs = re.split(r"\n\s*\n", source)
    return ParsedDocument(title="", blocks=[Block(_clean(p)) for p in paragraphs])


def _clean(text: str) -> str:
    return _SPACES_RE.sub(" ", text).strip()
